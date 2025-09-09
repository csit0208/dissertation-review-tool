import os
import re
from openai import OpenAI
from dotenv import load_dotenv

import pdfplumber
import docx
from sentence_transformers import SentenceTransformer, util
import requests
from bs4 import BeautifulSoup
from datetime import datetime
import csv
from rapidfuzz import fuzz
from fpdf import FPDF

load_dotenv()
client = OpenAI()

model = SentenceTransformer('all-MiniLM-L6-v2')

REFERENCE_TOPICS = {
    "public administration": "Dissertations in public administration should analyze governance, public policy, administrative law, ethical leadership, and organizational behavior in public sector contexts.",
    "applied psychology": "Applied psychology dissertations must explore practical applications of psychological theories, including behavioral interventions, assessment techniques, and workplace or clinical solutions.",
    "behavior analysis": "Behavior analysis research must involve measurable behavior change, reinforcement strategies, functional assessments, and empirical validation of interventions.",
    "business": "Business dissertations should cover areas such as strategic management, marketing theory, financial analysis, operational efficiency, and leadership practices.",
    "counseling": "Counseling dissertations should address therapeutic modalities, client outcomes, multicultural competence, mental health diagnostics, and ethical practice.",
    "education": "Education research should investigate curriculum design, learning theory, educational equity, assessment models, and pedagogical effectiveness.",
    "human services": "Human services dissertations should evaluate program outcomes, service delivery systems, vulnerable populations, policy impacts, and community engagement.",
    "public service": "Dissertations in public service must explore civic engagement, nonprofit leadership, service delivery, policy implementation, and public ethics.",
    "psychology": "Psychology dissertations must include theory-based analysis of cognition, emotion, development, psychopathology, and empirically validated methods.",
    "information technology": "Information technology research should focus on systems design, cybersecurity, data management, human-computer interaction, and technology integration in organizations.",
    "emergency management": "Emergency management dissertations must analyze disaster preparedness, crisis communication, risk mitigation, recovery operations, and interagency coordination.",
    "clinical psychology": "Clinical psychology research must address diagnostic criteria, evidence-based treatment, clinical outcomes, psychopathology, and therapeutic relationships.",
    "higher education leadership": "Dissertations in higher education leadership should examine institutional governance, leadership development, gender and diversity in academic leadership, and systemic challenges in postsecondary settings. Research may use qualitative, critical, or phenomenological methods to explore how leaders navigate organizational change, equity, and policy in higher education, including for-profit institutions.",
    "nursing": "Nursing dissertations focus on clinical practice, patient outcomes, healthcare policy, and evidence-based interventions. Research may use qualitative or quantitative methods to explore patient care, health systems, and clinical education.",
    "public health": "Public health dissertations analyze population health, epidemiology, community interventions, health equity, and policy outcomes. Topics may include disease prevention, health education, or public health program evaluation.",
    "social work": "Social work dissertations should analyze issues such as client advocacy, mental health policy, child welfare, community engagement, and evidence-based interventions across diver
}

def extract_text_from_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        with pdfplumber.open(file_path) as pdf:
            return "\n".join([page.extract_text() or "" for page in pdf.pages])
    elif ext == ".docx":
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    else:
        return f"Unsupported file type: {ext}"

def check_grammar(text):
    prompt = f"""
You are an academic writing assistant. Read the following dissertation excerpt and identify writing quality based on these criteria:
- Clear academic tone
- Grammar, punctuation, and usage
- Logical paragraphing

Rate overall writing on a scale from 0.0 (poor) to 1.0 (excellent).
Also, list up to 5 brief grammar or clarity issues with context.

Text:
{text[:3000]}

Respond only in this format:
Score: <float from 0.0 to 1.0>
Issues:
- <issue 1>
- <issue 2>
- ...
"""
    try:
        response = client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3,
            max_tokens=400
        )
        reply = response.choices[0].message.content.strip()
        lines = reply.splitlines()
        score_line = next((line for line in lines if line.lower().startswith("score:")), None)
        score = float(score_line.split(":")[1].strip()) if score_line else 0.7
        issues = [line[2:].strip() for line in lines if line.startswith("-")]
        return score, len(issues), [{"message": i, "context": "LLM-generated"} for i in issues]
    except Exception as e:
        # fallback in case of API failure
        return 0.7, 0, []
     
def extract_intext_citations(text):
    pattern = r'\(([A-Z][a-zA-Z\-’]+),\s*(\d{4})(?:[a-z])?(?:,\s*p\.?\s*\d+)?\)'
    return re.findall(pattern, text)

def extract_reference_entries(text):
    pattern = r'([A-Z][a-zA-Z’\-]+),?\s*([A-Z])?\.?\s*\((\d{4})\)'
    return re.findall(pattern, text)

def normalize_author(name):
    return name.lower().replace("’s", "").replace("’", "").replace("'", "").replace("-", "").strip()

def advanced_validate_apa_citations(text):
    intext = extract_intext_citations(text)
    refs = extract_reference_entries(text)
    ref_names = [(normalize_author(name), year) for name, _, year in refs]
    intext_normalized = [(normalize_author(name), year) for name, year in intext]

    KNOWN_GROUP_AUTHORS = {"nea", "bls", "ace", "wbi", "investopedia", "innosight", "constitution"}

    unmatched_intext = []
    unmatched_refs = []

    for name, year in intext_normalized:
        if name in KNOWN_GROUP_AUTHORS:
            continue
        match_found = any(
            fuzz.partial_ratio(name, rname) > 85 and year == ryear
            for rname, ryear in ref_names
        )
        if not match_found:
            unmatched_intext.append((name, year))

    for rname, ryear in ref_names:
        if rname in KNOWN_GROUP_AUTHORS:
            continue
        match_found = any(
            fuzz.partial_ratio(rname, name) > 85 and ryear == year
            for name, year in intext_normalized
        )
        if not match_found:
            unmatched_refs.append((rname, ryear))

    score = 1.0
    issues = []
    if unmatched_intext:
        issues.append(f"In-text citations not in reference list: {set(unmatched_intext)}")
        score -= 0.25
    if unmatched_refs:
        issues.append(f"Reference entries not cited in text: {set(unmatched_refs)}")
        score -= 0.25
    score = max(0.0, score)
    return round(score, 2), issues

def evaluate_with_llm(text, discipline):
    prompt = f"""
You are evaluating a doctoral dissertation in the field of {discipline}.

This program accepts both qualitative and quantitative research, including narrative, thematic, and interpretive studies. You are asked to score two rubric dimensions:

1. **Content Alignment** — Does the manuscript clearly reflect a valid research problem within {discipline}? Does the research design and writing reflect accepted scholarly approaches in this field?

2. **Accuracy of Outcomes and Conclusions** — Do the results, findings, or themes (as applicable) appear to have been analyzed appropriately? Do the conclusions follow logically from the data or themes, even if not presented statistically?

Please use the following rubric scoring:
- MET — The criterion is satisfied at a graduate level
- UNMET — The criterion is not demonstrated clearly or fully

Below is an excerpt from the dissertation, including material from the beginning and end:

{text[:3000]} ... [trimmed] ... {text[-3000:]}

Respond only using this exact format:
Content Alignment: <MET or UNMET>
Accuracy of Outcomes and Conclusions: <MET or UNMET>
"""

    try:
        response = client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3,
            max_tokens=200
        )
        reply = response.choices[0].message.content

        result = {
            "Content Alignment": "UNMET",
            "Accuracy of Outcomes and Conclusions": "UNMET",
            "_DEBUG_LLM_Response": reply
        }

        for line in reply.strip().splitlines():
            if "Content Alignment" in line:
                result["Content Alignment"] = "MET" if "MET" in line.upper() else "UNMET"
            elif "Accuracy" in line:
                result["Accuracy of Outcomes and Conclusions"] = "MET" if "MET" in line.upper() else "UNMET"

        return result

    except Exception as e:
        return {
            "Content Alignment": "UNMET",
            "Accuracy of Outcomes and Conclusions": "UNMET",
            "_DEBUG_LLM_Error": str(e)
        }

def rubric_evaluation(text, grammar_errors, citation_issues, topic, project_type):
    rubric = {}
    # --- Structure recognition based on project type ---
normalized = text.lower()

if project_type.lower() == "capstone":
    expected_headings = [
        "overview of the project", "problem statement and purpose", "theoretical framework", "project context",
        "historical background", "synthesis of the scholarly literature", "synthesis of the practitioner literature",
        "alignment of the project with the literature", "project questions", "project design", "stakeholders",
        "participants", "target audience", "role of the researcher", "study protocol", "sample",
        "data collection", "ethical considerations", "data analysis", "outcomes and findings", "application and benefits",
        "implications", "recommendations for policy", "recommendations for practice", "recommendations for future work",
        "conclusion"
    ]
else:  # Dissertation headings
    expected_headings = [
        "background of the study", "problem statement", "research question", "study rationale",
        "significance of the study", "identified gap", "overview of the methodology",
        "organization of the remainder of the study", "theoretical framework", "review of the literature",
        "synthesis of the literature", "research opportunities", "knowledge gaps", "proposed research",
        "sampling", "sample size", "data collection", "data analysis", "ethical considerations",
        "description of the sample", "presentation of data", "summary of findings", "limitations",
        "implications for policy", "recommendations for future research", "findings in context",
        "conclusion"
    ]

matched = [h for h in expected_headings if h in normalized]
rubric["_DEBUG_StructureHits"] = matched
rubric["Organization and Synthesis"] = "MET" if len(matched) >= int(0.6 * len(expected_headings)) else "UNMET"

llm_scores = evaluate_with_llm(text, topic)
rubric.update(llm_scores)

    normalized = text.lower()

    structure_terms = [
        "introduction",
        "literature review",
        "review of the literature",
        "methodology",
        "results",
        "presentation of the data",
        "discussion",
        "conclusion",
        "recommendations",
        "references"
    ]

    structure_hits = []
    for section in structure_terms:
        pattern = re.sub(r"\s+", r"\\s+", section)
        if re.search(pattern, normalized):
            structure_hits.append(section)

    rubric["_DEBUG_StructureHits"] = structure_hits
    rubric["Organization and Synthesis"] = (
        "MET" if len(structure_hits) >= 4 else "UNMET"
    )

    avg_errors_per_page = grammar_errors / max(1, (len(text) // 250))
    critical_citation_issues = [
        issue for issue in citation_issues
        if "not in reference list" in issue or "not cited in text" in issue
    ]

    rubric["Writing Mechanics, APA, Citations, Evidence"] = (
        "MET" if avg_errors_per_page <= 5 and len(critical_citation_issues) <= 5 else "UNMET"
    )

    if rubric["Organization and Synthesis"] == "UNMET":
        rubric["Faculty Comment"] = "Consider clarifying or labeling section transitions to reflect APA structure more clearly (e.g., Introduction, Methodology, etc.)."
    elif rubric["Writing Mechanics, APA, Citations, Evidence"] == "UNMET":
        rubric["Faculty Comment"] = "Writing mechanics or citation formatting require moderate revision."
    elif rubric["Accuracy of Outcomes and Conclusions"] == "UNMET":
        rubric["Faculty Comment"] = "Ensure that conclusions are tightly tied to presented results, even in qualitative analysis."
    else:
        rubric["Faculty Comment"] = "Dissertation meets expectations for structure, alignment, conclusions, and APA standards."

    return rubric

def evaluate_dissertation(text, topic):
    grammar_score, grammar_errors, grammar_issues = check_grammar(text)
    citation_score, citation_issues = advanced_validate_apa_citations(text)
    rubric = rubric_evaluation(text, grammar_errors, citation_issues, topic)
    met_count = sum(1 for v in rubric.values() if v == "MET")
    decision = (
        "Approved" if met_count == 4 else
        "Needs Further Review" if met_count == 3 else
        "Not Approved"
    )
    overall = round((grammar_score * 0.3 + citation_score * 0.3 + (met_count / 4) * 0.4), 2)

    return {
        "decision": decision,
        "scores": {
            "grammar": round(grammar_score, 2),
            "citations": round(citation_score, 2),
            "overall": overall
        },
        "rubric": rubric,
        "details": {
            "grammar_errors": grammar_errors,
            "grammar_issues": grammar_issues,
            "citation_issues": citation_issues
        }
    }

def export_to_csv(file_path, topic, result, reviewer="AutoReviewer", output_path="dissertation_reviews.csv"):
    headers = [
        "timestamp", "reviewer", "file", "discipline", "decision",
        "grammar", "citations", "overall",
        "grammar_errors", "citation_issue_count"
    ] + [key for key in result["rubric"] if not key.startswith("_DEBUG")]

    row = [
        datetime.now().isoformat(),
        reviewer,
        os.path.basename(file_path),
        topic,
        result["decision"],
        result["scores"]["grammar"],
        result["scores"]["citations"],
        result["scores"]["overall"],
        result["details"]["grammar_errors"],
        len(result["details"]["citation_issues"])
    ] + [result["rubric"][key] for key in result["rubric"] if not key.startswith("_DEBUG")]

    file_exists = os.path.exists(output_path)
    with open(output_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(headers)
        writer.writerow(row)

def export_to_pdf(file_path, topic, result, reviewer="AutoReviewer", output_dir="."):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Dissertation Evaluation Report", ln=1, align="C")
    pdf.ln(10)

    pdf.cell(200, 10, txt=f"Reviewer: {reviewer}", ln=1)
    pdf.cell(200, 10, txt=f"File: {os.path.basename(file_path)}", ln=1)
    pdf.cell(200, 10, txt=f"Discipline: {topic}", ln=1)
    pdf.cell(200, 10, txt=f"Decision: {result['decision']}", ln=1)
    pdf.ln(5)

    pdf.set_font("Arial", style="B", size=12)
    pdf.cell(200, 10, txt="Scores:", ln=1)
    pdf.set_font("Arial", size=12)
    for key, val in result["scores"].items():
        pdf.cell(200, 10, txt=f"{key.title()}: {val}", ln=1)

    pdf.ln(5)
    pdf.set_font("Arial", style="B", size=12)
    pdf.cell(200, 10, txt="Rubric Ratings:", ln=1)
    pdf.set_font("Arial", size=12)
    for key, val in result["rubric"].items():
        if not key.startswith("_DEBUG"):
            pdf.cell(200, 10, txt=f"{key}: {val}", ln=1)

    filename = f"evaluation_{os.path.basename(file_path).split('.')[0]}.pdf"
    pdf.output(os.path.join(output_dir, filename))
